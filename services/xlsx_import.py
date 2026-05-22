"""One-time importer: reads side.xlsx and populates data/*.json files."""
import re
import pathlib
from helpers import load, save


def slugify(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", name.strip().lower()).strip("-")


def col_index(col: str) -> int:
    """'A'->0, 'B'->1, 'AA'->26, etc."""
    result = 0
    for ch in col.upper():
        result = result * 26 + (ord(ch) - ord("A") + 1)
    return result - 1


def import_from_xlsx(path: str = "side.xlsx") -> dict:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)

    all_opponent_decks: dict[str, dict] = {}
    my_decks_list = []
    sideboard_plans: dict[str, dict] = {}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows or all(v is None for v in rows[0]):
            continue

        # --- Row 0 (Row 1 in Excel): category headers ---
        cat_row = rows[0] if rows else []
        # col -> category: categories appear at specific column indices
        col_to_cat: dict[int, str] = {}
        current_cat = "Aggro"
        for ci, val in enumerate(cat_row):
            if val and str(val).strip().lower() in ("aggro", "midrange", "control", "combo"):
                current_cat = str(val).strip().capitalize()
            col_to_cat[ci] = current_cat

        # --- Row 1 (Row 2 in Excel): opponent deck names from col 2 onward ---
        if len(rows) < 2:
            continue
        header_row = rows[1]
        # col_index -> opponent_deck_id
        col_to_opponent: dict[int, str] = {}
        for ci, val in enumerate(header_row):
            if ci < 2 or not val or str(val).strip() in ("#", "Main", "Side"):
                continue
            name = str(val).strip()
            opp_id = f"import-{slugify(name)}"
            col_to_opponent[ci] = opp_id
            cat = col_to_cat.get(ci, "Aggro")
            if opp_id not in all_opponent_decks:
                all_opponent_decks[opp_id] = {
                    "id": opp_id,
                    "name": name,
                    "category": cat,
                    "meta_share": 0.0,
                    "archetype_id": None,
                    "top_decklist_url": None,
                    "included": True,
                }

        # --- Find divider row (A = "Side" or "SIDE") ---
        # Use the FIRST row where A == "side"; the B column value (often a total count) is irrelevant.
        divider_row_idx = None
        for ri, row in enumerate(rows[2:], start=2):
            cell_a = row[0]
            if cell_a and str(cell_a).strip().lower() == "side":
                divider_row_idx = ri
                break

        if divider_row_idx is None:
            continue

        # --- Parse main deck rows ---
        main_deck_data: list[dict] = []
        out_by_card: dict[str, dict[str, int]] = {}

        for row in rows[2:divider_row_idx]:
            card_name = row[0]
            qty_raw = row[1] if len(row) > 1 else None
            if not card_name or str(card_name).strip() in ("", "Main", "Side", "#"):
                continue
            card_name = str(card_name).strip()
            try:
                qty = int(qty_raw) if qty_raw else 0
            except (ValueError, TypeError):
                qty = 0
            if qty == 0 and not any(
                row[ci] for ci in col_to_opponent if ci < len(row)
            ):
                continue

            main_deck_data.append({"name": card_name, "quantity": qty})
            out_by_card[card_name] = {}
            for ci, opp_id in col_to_opponent.items():
                if ci < len(row) and row[ci]:
                    try:
                        v = int(row[ci])
                        if v > 0:
                            out_by_card[card_name][opp_id] = v
                    except (ValueError, TypeError):
                        pass

        # --- Parse sideboard rows ---
        sideboard_data: list[dict] = []
        in_by_card: dict[str, dict[str, int]] = {}

        for row in rows[divider_row_idx + 1:]:
            card_name = row[0]
            qty_raw = row[1] if len(row) > 1 else None
            if not card_name or str(card_name).strip() in ("", "Main", "Side", "#"):
                continue
            card_name = str(card_name).strip()
            try:
                qty = int(qty_raw) if qty_raw else 0
            except (ValueError, TypeError):
                qty = 0
            # skip counter rows (all numeric, large totals)
            if all(isinstance(v, (int, float)) or v is None for v in row) and qty > 10:
                continue

            sideboard_data.append({"name": card_name, "quantity": qty})
            in_by_card[card_name] = {}
            for ci, opp_id in col_to_opponent.items():
                if ci < len(row) and row[ci]:
                    try:
                        v = int(row[ci])
                        if v > 0:
                            in_by_card[card_name][opp_id] = v
                    except (ValueError, TypeError):
                        pass

        # --- Build sideboard plans for this deck ---
        deck_id = f"import-{slugify(sheet_name)}"
        plans_for_deck: dict[str, dict] = {}
        for opp_id in col_to_opponent.values():
            out_entries = [
                {"card": cn, "quantity": q}
                for cn, opp_map in out_by_card.items()
                if opp_id in opp_map
                for q in [opp_map[opp_id]]
            ]
            in_entries = [
                {"card": cn, "quantity": q}
                for cn, opp_map in in_by_card.items()
                if opp_id in opp_map
                for q in [opp_map[opp_id]]
            ]
            if out_entries or in_entries:
                plans_for_deck[opp_id] = {
                    "out": out_entries,
                    "in": in_entries,
                    "notes": "",
                }

        sideboard_plans[deck_id] = plans_for_deck
        my_decks_list.append({
            "id": deck_id,
            "archidekt_id": None,
            "name": sheet_name,
            "archidekt_url": None,
            "last_synced": None,
            "main_deck": main_deck_data,
            "sideboard": sideboard_data,
        })

    wb.close()

    # --- Save meta_decks.json (preserve existing user overrides) ---
    existing_meta = load("meta_decks.json")
    existing_by_id = {d["id"]: d for d in existing_meta.get("decks", [])}
    merged_decks = []
    for opp_id, deck in all_opponent_decks.items():
        if opp_id in existing_by_id:
            existing = existing_by_id[opp_id]
            deck["category"] = existing.get("category", deck["category"])
            deck["included"] = existing.get("included", True)
        merged_decks.append(deck)

    save("meta_decks.json", {
        "last_fetched": existing_meta.get("last_fetched"),
        "meta_id": existing_meta.get("meta_id"),
        "decks": merged_decks,
    })

    # --- Save my_decks.json (merge with existing Archidekt-synced data) ---
    existing_my = load("my_decks.json")
    existing_my_by_name = {
        d["name"].lower(): d for d in existing_my.get("decks", [])
    }
    merged_my_decks = []
    for new_deck in my_decks_list:
        existing = existing_my_by_name.get(new_deck["name"].lower())
        if existing and existing.get("archidekt_id"):
            # Archidekt data takes precedence; only import if no sync yet
            merged_my_decks.append(existing)
        else:
            merged_my_decks.append(new_deck)

    save("my_decks.json", {
        "folder_url": existing_my.get("folder_url", "https://archidekt.com/folders/1123156"),
        "last_folder_sync": existing_my.get("last_folder_sync"),
        "decks": merged_my_decks,
    })

    # --- Save sideboards.json (merge, don't overwrite manual edits) ---
    existing_sb = load("sideboards.json")
    existing_plans = existing_sb.get("plans", {})
    for deck_id, plans in sideboard_plans.items():
        if deck_id not in existing_plans:
            existing_plans[deck_id] = plans

    save("sideboards.json", {
        "plans": existing_plans,
        "theoretical_pool": existing_sb.get("theoretical_pool", {}),
    })

    return {
        "imported_decks": len(my_decks_list),
        "imported_opponent_decks": len(all_opponent_decks),
        "deck_names": [d["name"] for d in my_decks_list],
    }
